import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import os
import csv
import shutil

# Import sample orders from manifest generator
from generate_delivery_manifest import orders_sample

def style_header(cell):
    cell.fill = PatternFill(start_color="111111", end_color="111111", fill_type="solid")
    cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style='thin', color='333333'),
        right=Side(style='thin', color='333333'),
        top=Side(style='medium', color='111111'),
        bottom=Side(style='medium', color='111111')
    )

thin_border = Border(
    left=Side(style='thin', color='D8D8D8'),
    right=Side(style='thin', color='D8D8D8'),
    top=Side(style='thin', color='D8D8D8'),
    bottom=Side(style='thin', color='D8D8D8')
)

def create_manual_template_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Suivi Livraisons Santos"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        ("A", "Tracking_Number", 18, "SN-TRK-001"),
        ("B", "Order_ID", 14, "SN-1001"),
        ("C", "Date", 14, "2026-10-30"),
        ("D", "Customer_Name", 24, "Exemple: Amina Tahiri"),
        ("E", "Phone_1", 16, "0661000000"),
        ("F", "Phone_2_WhatsApp", 18, "0661000000"),
        ("G", "City", 16, "Casablanca"),
        ("H", "District", 18, "Gauthier"),
        ("I", "Full_Address", 35, "12 Rue Jean Jaurès Apt 4"),
        ("J", "Product_Items", 32, "MANTECO ORIGINALLY"),
        ("K", "Size", 10, "M"),
        ("L", "Color", 12, "Noir"),
        ("M", "Quantity", 10, 1),
        ("N", "COD_Amount_MAD", 18, 850.00),
        ("O", "Carrier", 20, "Cathedis"),
        ("P", "Confirmation_Status", 22, "Confirmé WhatsApp"),
        ("Q", "Delivery_Status", 22, "À expédier"),
        ("R", "Payment_Status", 20, "À encaisser"),
        ("S", "Courier_Notes", 35, "Appeler 30min avant passage"),
    ]
    
    sample_row_2 = [
        "SN-TRK-002", "SN-1002", "2026-10-30", "Exemple: Mehdi Berrada",
        "0662000000", "0662000000", "Rabat", "Agdal",
        "45 Avenue de France", "Trensh original", "L", "Camel",
        1, 700.00, "Amana", "Confirmé Appel", "À expédier", "À encaisser",
        "Livraison bureau après 14h"
    ]
    
    ws.row_dimensions[1].height = 32
    
    for col_idx, (col_letter, header_name, width, _) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_name)
        style_header(cell)
        ws.column_dimensions[col_letter].width = width

    sample_rows = [
        [h[3] for h in headers],
        sample_row_2
    ]
    
    sample_fill = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")
    sample_font = Font(name="Segoe UI", size=10, italic=True, color="333333")
    regular_font = Font(name="Segoe UI", size=10, color="111111")

    for r_idx, row_data in enumerate(sample_rows, start=2):
        ws.row_dimensions[r_idx].height = 24
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = sample_font
            cell.fill = sample_fill
            cell.border = thin_border
            
            if c_idx in [5, 6]:
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [1, 2, 3, 11, 12, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 14:
                cell.number_format = '#,##0.00 "DH"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [15, 16, 17, 18]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # 50 blank formatted rows for manual data entry
    for r_idx in range(4, 54):
        ws.row_dimensions[r_idx].height = 22
        is_even = (r_idx % 2 == 0)
        row_fill = PatternFill(start_color="FFFFFF" if is_even else "FAFAFA", end_color="FFFFFF" if is_even else "FAFAFA", fill_type="solid")
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx, value="")
            cell.font = regular_font
            cell.fill = row_fill
            cell.border = thin_border
            if c_idx in [5, 6]:
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [1, 2, 3, 11, 12, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 14:
                cell.number_format = '#,##0.00 "DH"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [15, 16, 17, 18]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Dropdowns for clean manual entry
    dv_carrier = DataValidation(type="list", formula1='"Cathedis,Amana,Sendit,Livreur Privé Santos,Autre"', allow_blank=True)
    ws.add_data_validation(dv_carrier)
    dv_carrier.add("O2:O53")

    dv_conf = DataValidation(type="list", formula1='"Confirmé WhatsApp,Confirmé Appel,Pas de réponse,Injoignable,Annulé,À rappeler"', allow_blank=True)
    ws.add_data_validation(dv_conf)
    dv_conf.add("P2:P53")

    dv_deliv = DataValidation(type="list", formula1='"À expédier,En cours de préparation,En cours de livraison,Livré,Retourné,Refusé,Reporté"', allow_blank=True)
    ws.add_data_validation(dv_deliv)
    dv_deliv.add("Q2:Q53")

    dv_pay = DataValidation(type="list", formula1='"À encaisser,Encaissé,Virement reçu,Remboursé"', allow_blank=True)
    ws.add_data_validation(dv_pay)
    dv_pay.add("R2:R53")

    # Freeze top row
    ws.freeze_panes = "A2"

    wb.save("santos_delivery_manual_template.xlsx")
    shutil.copyfile("santos_delivery_manual_template.xlsx", "santos_delivery_manual_template_excel.xlsx")
    print("Saved santos_delivery_manual_template.xlsx and santos_delivery_manual_template_excel.xlsx")

def create_tracking_orders_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Suivi des Commandes Santos"
    ws.views.sheetView[0].showGridLines = True

    columns = [
        ("A", "Tracking_Number", 18),
        ("B", "Order_ID", 14),
        ("C", "Date", 18),
        ("D", "Customer_Name", 24),
        ("E", "Phone", 16),
        ("F", "City", 16),
        ("G", "Address", 40),
        ("H", "Product", 36),
        ("I", "Quantity", 10),
        ("J", "COD_Amount_MAD", 18),
        ("K", "Carrier", 22),
        ("L", "Delivery_Status", 24),
        ("M", "Payment_Status", 24),
        ("N", "Last_Update", 18),
        ("O", "Notes", 38),
    ]

    ws.row_dimensions[1].height = 32
    for col_idx, (col_letter, col_name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        style_header(cell)
        ws.column_dimensions[col_letter].width = width

    font_data = Font(name="Segoe UI", size=10, color="111111")
    fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_odd = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for r_idx, order in enumerate(orders_sample, start=2):
        ws.row_dimensions[r_idx].height = 24
        row_fill = fill_even if r_idx % 2 == 0 else fill_odd
        
        row_vals = [
            order['Tracking_Number'],
            order['Order_ID'],
            order['Date'],
            order['Customer_Name'],
            order['Phone'],
            order['City'],
            order['Address'],
            order['Product'],
            order['Quantity'],
            order['COD_Amount_MAD'],
            order['Carrier'],
            order['Delivery_Status'],
            order['Payment_Status'],
            order['Last_Update'],
            order['Notes']
        ]

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = thin_border
            
            # Formatting
            if c_idx == 5: # Phone
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [1, 2, 3, 9, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 10: # COD Amount
                cell.number_format = '#,##0.00 "DH"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [6, 11, 12, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.freeze_panes = "A2"
    wb.save("santos_delivery_tracking.xlsx")
    print("Saved santos_delivery_tracking.xlsx")

def fix_csv_files():
    # Headers
    headers = [
        "Tracking_Number", "Order_ID", "Date", "Customer_Name", "Phone_1", "Phone_2_WhatsApp",
        "City", "District", "Full_Address", "Product_Items", "Size", "Color", "Quantity",
        "COD_Amount_MAD", "Carrier", "Confirmation_Status", "Delivery_Status", "Payment_Status", "Courier_Notes"
    ]
    sample_1 = [
        "SN-TRK-001", "SN-1001", "2026-10-30", "Exemple: Amina Tahiri", "0661000000", "0661000000",
        "Casablanca", "Gauthier", "12 Rue Jean Jaurès Apt 4", "MANTECO ORIGINALLY", "M", "Noir", 1,
        "850.00", "Cathedis", "Confirmé WhatsApp", "À expédier", "À encaisser", "Appeler 30min avant passage"
    ]
    sample_2 = [
        "SN-TRK-002", "SN-1002", "2026-10-30", "Exemple: Mehdi Berrada", "0662000000", "0662000000",
        "Rabat", "Agdal", "45 Avenue de France", "Trensh original", "L", "Camel", 1,
        "700.00", "Amana", "Confirmé Appel", "À expédier", "À encaisser", "Livraison bureau après 14h"
    ]

    # Clean standard comma CSV with UTF-8 BOM
    with open("santos_delivery_manual_template.csv", "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=",", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        writer.writerow(sample_1)
        writer.writerow(sample_2)
    print("Saved clean santos_delivery_manual_template.csv")

    with open("santos_delivery_manual_template_excel.csv", "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=",", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        writer.writerow(sample_1)
        writer.writerow(sample_2)
    print("Saved clean santos_delivery_manual_template_excel.csv")

    # Also update santos_delivery_tracking_excel.csv with comma so WPS Office never puts everything in Col A
    tracking_headers = [
        'Tracking_Number', 'Order_ID', 'Date', 'Customer_Name', 'Phone',
        'City', 'Address', 'Product', 'Quantity', 'COD_Amount_MAD',
        'Carrier', 'Delivery_Status', 'Payment_Status', 'Last_Update', 'Notes'
    ]
    with open('santos_delivery_tracking_excel.csv', 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=tracking_headers, delimiter=',')
        writer.writeheader()
        writer.writerows(orders_sample)
    print("Saved santos_delivery_tracking_excel.csv with standard comma delimiter")

if __name__ == "__main__":
    create_manual_template_xlsx()
    create_tracking_orders_xlsx()
    fix_csv_files()
